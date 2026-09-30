from __future__ import annotations

from io import BytesIO
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from formatting.sanitize import (
    parse_terms,
    sanitize_text,
)


def set_cell_shading(
    cell,
    fill: str = "1F2937",
) -> None:

    tc_pr = cell._tc.get_or_add_tcPr()

    shading = OxmlElement(
        "w:shd"
    )

    shading.set(
        qn("w:fill"),
        fill,
    )

    tc_pr.append(
        shading
    )


def set_default_font(
    document: Document,
) -> None:

    normal = document.styles[
        "Normal"
    ]

    normal.font.name = (
        "Times New Roman"
    )

    normal.font.size = Pt(11)

    normal._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Times New Roman",
    )


def add_footer(
    document: Document,
    company_name: str,
) -> None:

    for section in document.sections:

        footer = section.footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            f"{company_name} | "
            "AI drafting assistance - "
            "review before use"
        )

        run.font.size = Pt(8)


def add_logo(
    document: Document,
    logo_path: str | None = None,
    logo_bytes: bytes | None = None,
) -> None:

    if not logo_path and not logo_bytes:
        return

    if logo_bytes:

        source = BytesIO(
            logo_bytes
        )

    else:

        source = logo_path

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    paragraph.add_run().add_picture(
        source,
        width=Inches(1.1),
    )


def add_signature_table(
    document: Document,
) -> None:

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.style = "Table Grid"

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.cell(
        0,
        0,
    ).text = (
        "Party 1: __________________________"
    )

    table.cell(
        0,
        1,
    ).text = (
        "Party 2: __________________________"
    )

    table.cell(
        1,
        0,
    ).text = (
        "Date: _____________________________"
    )

    table.cell(
        1,
        1,
    ).text = (
        "Date: _____________________________"
    )


def format_docx(
    text: str,
    doc_type: str,
    *,
    terms: str = "",
    logo_path: str | None = None,
    logo_bytes: bytes | None = None,
    company_name: str = "LegalEase",
) -> bytes:

    document = Document()

    set_default_font(
        document
    )

    for section in document.sections:

        section.top_margin = Inches(
            0.7
        )

        section.bottom_margin = Inches(
            0.7
        )

        section.left_margin = Inches(
            0.85
        )

        section.right_margin = Inches(
            0.85
        )

    add_footer(
        document,
        company_name,
    )

    add_logo(
        document,
        logo_path=logo_path,
        logo_bytes=logo_bytes,
    )

    title = document.add_heading(
        doc_type.upper(),
        level=0,
    )

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    clean_text = sanitize_text(
        text
    )

    for raw_line in clean_text.splitlines():

        line = raw_line.strip()

        if not line:
            continue

        if line.startswith("# "):

            document.add_heading(
                line[2:].strip(),
                level=1,
            )

        elif line.startswith("## "):

            document.add_heading(
                line[3:].strip(),
                level=2,
            )

        elif line.startswith("### "):

            document.add_heading(
                line[4:].strip(),
                level=3,
            )

        elif line.startswith(
            ("- ", "* ")
        ):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                line[2:].strip()
            )

        elif re.match(
            r"^\d+[.)]\s+",
            line,
        ):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                line
            )

            run.bold = True

        elif line.startswith("> "):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                line[2:].strip()
            )

            run.italic = True

        else:

            paragraph = document.add_paragraph(
                line
            )

            paragraph.paragraph_format.space_after = Pt(
                5
            )

    term_items = parse_terms(
        terms
    )

    if term_items:

        document.add_heading(
            "Key Terms",
            level=2,
        )

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.style = "Table Grid"

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        header = table.rows[
            0
        ].cells

        header[0].text = "#"
        header[1].text = "Term"

        set_cell_shading(
            header[0]
        )

        set_cell_shading(
            header[1]
        )

        for index, term in enumerate(
            term_items,
            start=1,
        ):

            row = table.add_row().cells

            row[0].text = str(
                index
            )

            row[1].text = term

    document.add_heading(
        "Signatures",
        level=2,
    )

    add_signature_table(
        document
    )

    note = document.add_paragraph()

    run = note.add_run(
        "Drafting notice: AI-generated "
        "content is not a substitute for "
        "advice from a qualified lawyer "
        "licensed in the relevant jurisdiction."
    )

    run.italic = True

    run.font.size = Pt(8)

    output = BytesIO()

    document.save(
        output
    )

    return output.getvalue()